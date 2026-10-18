import copy, re, sys, os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SP = os.path.dirname(os.path.abspath(__file__))
FIGS = r"S:\Projects\Urban Farm\paper\figures"
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(SP, "paper_v2_raw.docx")

# =========================================================================== references
REFS = {
 "un2018": "United Nations, Department of Economic and Social Affairs, Population Division, \u201cWorld urbanization prospects: The 2018 revision,\u201d United Nations, New York, NY, USA, Rep. ST/ESA/SER.A/420, 2019.",
 "fao2017": "Food and Agriculture Organization of the United Nations, The Future of Food and Agriculture\u2014Trends and Challenges. Rome, Italy: FAO, 2017.",
 "despommier": "D. Despommier, The Vertical Farm: Feeding the World in the 21st Century. New York, NY, USA: Thomas Dunne Books, 2010.",
 "kozai": "T. Kozai, G. Niu, and M. Takagaki, Eds., Plant Factory: An Indoor Vertical Farming System for Efficient Quality Food Production, 2nd ed. London, U.K.: Academic Press, 2019.",
 "benke": "K. Benke and B. Tomkins, \u201cFuture food-production systems: Vertical farming and controlled-environment agriculture,\u201d Sustainability: Sci., Pract. Policy, vol. 13, no. 1, pp. 13\u201326, 2017.",
 "savary": "S. Savary, L. Willocquet, S. J. Pethybridge, P. Esker, N. McRoberts, and A. Nelson, \u201cThe global burden of pathogens and pests on major food crops,\u201d Nature Ecol. Evol., vol. 3, no. 3, pp. 430\u2013439, 2019.",
 "liakos": "K. G. Liakos, P. Busato, D. Moshou, S. Pearson, and D. Bochtis, \u201cMachine learning in agriculture: A review,\u201d Sensors, vol. 18, no. 8, Art. no. 2674, 2018.",
 "kamilaris": "A. Kamilaris and F. X. Prenafeta-Bold\u00fa, \u201cDeep learning in agriculture: A survey,\u201d Comput. Electron. Agric., vol. 147, pp. 70\u201390, 2018.",
 "mohanty": "S. P. Mohanty, D. P. Hughes, and M. Salath\u00e9, \u201cUsing deep learning for image-based plant disease detection,\u201d Front. Plant Sci., vol. 7, Art. no. 1419, 2016.",
 "ferentinos": "K. P. Ferentinos, \u201cDeep learning models for plant disease detection and diagnosis,\u201d Comput. Electron. Agric., vol. 145, pp. 311\u2013318, 2018.",
 "wolfert": "S. Wolfert, L. Ge, C. Verdouw, and M.-J. Bogaardt, \u201cBig data in smart farming\u2014A review,\u201d Agric. Syst., vol. 153, pp. 69\u201380, 2017.",
 "elijah": "O. Elijah, T. A. Rahman, I. Orikumhi, C. Y. Leow, and M. N. Hindia, \u201cAn overview of Internet of Things (IoT) and data analytics in agriculture: Benefits and challenges,\u201d IEEE Internet Things J., vol. 5, no. 5, pp. 3758\u20133773, Oct. 2018.",
 "lewis": "P. Lewis et al., \u201cRetrieval-augmented generation for knowledge-intensive NLP tasks,\u201d in Proc. Adv. Neural Inf. Process. Syst. (NeurIPS), 2020, pp. 9459\u20139474.",
 "krizhevsky": "A. Krizhevsky, I. Sutskever, and G. E. Hinton, \u201cImageNet classification with deep convolutional neural networks,\u201d in Proc. Adv. Neural Inf. Process. Syst. (NeurIPS), 2012, pp. 1097\u20131105.",
 "deng": "J. Deng, W. Dong, R. Socher, L.-J. Li, K. Li, and L. Fei-Fei, \u201cImageNet: A large-scale hierarchical image database,\u201d in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), 2009, pp. 248\u2013255.",
 "lecun": "Y. LeCun, Y. Bengio, and G. Hinton, \u201cDeep learning,\u201d Nature, vol. 521, no. 7553, pp. 436\u2013444, 2015.",
 "hughes": "D. P. Hughes and M. Salath\u00e9, \u201cAn open access repository of images on plant health to enable the development of mobile disease diagnostics,\u201d arXiv:1511.08060, 2015.",
 "sladojevic": "S. Sladojevic, M. Arsenovic, A. Anderla, D. Culibrk, and D. Stefanovic, \u201cDeep neural networks based recognition of plant diseases by leaf image classification,\u201d Comput. Intell. Neurosci., vol. 2016, Art. no. 3289801, 2016.",
 "brahimi": "M. Brahimi, K. Boukhalfa, and A. Moussaoui, \u201cDeep learning for tomato diseases: Classification and symptoms visualization,\u201d Appl. Artif. Intell., vol. 31, no. 4, pp. 299\u2013315, 2017.",
 "too": "E. C. Too, L. Yujian, S. Njuki, and L. Yingchun, \u201cA comparative study of fine-tuning deep learning models for plant disease identification,\u201d Comput. Electron. Agric., vol. 161, pp. 272\u2013279, 2019.",
 "he": "K. He, X. Zhang, S. Ren, and J. Sun, \u201cDeep residual learning for image recognition,\u201d in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), 2016, pp. 770\u2013778.",
 "thapa": "R. Thapa, K. Zhang, N. Snavely, S. Belongie, and A. Khan, \u201cThe Plant Pathology Challenge 2020 data set to classify foliar disease of apples,\u201d Appl. Plant Sci., vol. 8, no. 9, Art. no. e11390, 2020.",
 "vaswani": "A. Vaswani et al., \u201cAttention is all you need,\u201d in Proc. Adv. Neural Inf. Process. Syst. (NeurIPS), 2017, pp. 5998\u20136008.",
 "dosovitskiy": "A. Dosovitskiy et al., \u201cAn image is worth 16x16 words: Transformers for image recognition at scale,\u201d in Proc. Int. Conf. Learn. Represent. (ICLR), 2021.",
 "touvron": "H. Touvron, M. Cord, M. Douze, F. Massa, A. Sablayrolles, and H. J\u00e9gou, \u201cTraining data-efficient image transformers & distillation through attention,\u201d in Proc. Int. Conf. Mach. Learn. (ICML), 2021, pp. 10347\u201310357.",
 "borhani": "Y. Borhani, J. Khoramdel, and E. Najafi, \u201cA deep learning based approach for automated plant disease classification using vision transformer,\u201d Sci. Rep., vol. 12, Art. no. 11554, 2022.",
 "breiman": "L. Breiman, \u201cRandom forests,\u201d Mach. Learn., vol. 45, no. 1, pp. 5\u201332, 2001.",
 "hoerl": "A. E. Hoerl and R. W. Kennard, \u201cRidge regression: Biased estimation for nonorthogonal problems,\u201d Technometrics, vol. 12, no. 1, pp. 55\u201367, 1970.",
 "shamshiri": "R. R. Shamshiri et al., \u201cAdvances in greenhouse automation and controlled environment agriculture: A transition to plant factories and urban agriculture,\u201d Int. J. Agric. Biol. Eng., vol. 11, no. 1, pp. 1\u201322, 2018.",
 "hochreiter": "S. Hochreiter and J. Schmidhuber, \u201cLong short-term memory,\u201d Neural Comput., vol. 9, no. 8, pp. 1735\u20131780, 1997.",
 "hyndman": "R. J. Hyndman and G. Athanasopoulos, Forecasting: Principles and Practice, 3rd ed. Melbourne, VIC, Australia: OTexts, 2021.",
 "tetens": "O. Tetens, \u201c\u00dcber einige meteorologische Begriffe,\u201d Z. Geophys., vol. 6, pp. 297\u2013309, 1930.",
 "monteith": "J. L. Monteith and M. H. Unsworth, Principles of Environmental Physics, 4th ed. Oxford, U.K.: Academic Press, 2013.",
 "liebig": "J. von Liebig, Die organische Chemie in ihrer Anwendung auf Agricultur und Physiologie. Braunschweig, Germany: Vieweg, 1840.",
 "devlin": "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, \u201cBERT: Pre-training of deep bidirectional transformers for language understanding,\u201d in Proc. Conf. North Amer. Chapter Assoc. Comput. Linguistics (NAACL-HLT), 2019, pp. 4171\u20134186.",
 "reimers": "N. Reimers and I. Gurevych, \u201cSentence-BERT: Sentence embeddings using Siamese BERT-networks,\u201d in Proc. Conf. Empirical Methods Natural Lang. Process. (EMNLP-IJCNLP), 2019, pp. 3982\u20133992.",
 "wang": "W. Wang, F. Wei, L. Dong, H. Bao, N. Yang, and M. Zhou, \u201cMiniLM: Deep self-attention distillation for task-agnostic compression of pre-trained transformers,\u201d in Proc. Adv. Neural Inf. Process. Syst. (NeurIPS), 2020, pp. 5776\u20135788.",
 "salton": "G. Salton and C. Buckley, \u201cTerm-weighting approaches in automatic text retrieval,\u201d Inf. Process. Manage., vol. 24, no. 5, pp. 513\u2013523, 1988.",
 "robertson": "S. Robertson and H. Zaragoza, \u201cThe probabilistic relevance framework: BM25 and beyond,\u201d Found. Trends Inf. Retr., vol. 3, no. 4, pp. 333\u2013389, 2009.",
 "karpukhin": "V. Karpukhin et al., \u201cDense passage retrieval for open-domain question answering,\u201d in Proc. Conf. Empirical Methods Natural Lang. Process. (EMNLP), 2020, pp. 6769\u20136781.",
 "gemma": "Gemma Team, \u201cGemma 3 technical report,\u201d arXiv:2503.19786, 2025.",
 "nikolenko": "S. I. Nikolenko, Synthetic Data for Deep Learning. Cham, Switzerland: Springer, 2021.",
 "wolf": "T. Wolf et al., \u201cTransformers: State-of-the-art natural language processing,\u201d in Proc. Conf. Empirical Methods Natural Lang. Process.: Syst. Demonstrations (EMNLP), 2020, pp. 38\u201345.",
 "hfdisease": "wambugu71, \u201ccrop_leaf_diseases_vit,\u201d Hugging Face Model Hub. [Online]. Available: https://huggingface.co/wambugu71/crop_leaf_diseases_vit",
 "hfspecies": "dima806, \u201cmedicinal_plants_image_detection,\u201d Hugging Face Model Hub. [Online]. Available: https://huggingface.co/dima806/medicinal_plants_image_detection",
 "paszke": "A. Paszke et al., \u201cPyTorch: An imperative style, high-performance deep learning library,\u201d in Proc. Adv. Neural Inf. Process. Syst. (NeurIPS), 2019, pp. 8024\u20138035.",
 "ioffe": "S. Ioffe and C. Szegedy, \u201cBatch normalization: Accelerating deep network training by reducing internal covariate shift,\u201d in Proc. Int. Conf. Mach. Learn. (ICML), 2015, pp. 448\u2013456.",
 "srivastava": "N. Srivastava, G. Hinton, A. Krizhevsky, I. Sutskever, and R. Salakhutdinov, \u201cDropout: A simple way to prevent neural networks from overfitting,\u201d J. Mach. Learn. Res., vol. 15, no. 56, pp. 1929\u20131958, 2014.",
 "shorten": "C. Shorten and T. M. Khoshgoftaar, \u201cA survey on image data augmentation for deep learning,\u201d J. Big Data, vol. 6, Art. no. 60, 2019.",
 "kingma": "D. P. Kingma and J. Ba, \u201cAdam: A method for stochastic optimization,\u201d in Proc. Int. Conf. Learn. Represent. (ICLR), 2015.",
 "pedregosa": "F. Pedregosa et al., \u201cScikit-learn: Machine learning in Python,\u201d J. Mach. Learn. Res., vol. 12, pp. 2825\u20132830, 2011.",
 "tibshirani": "R. Tibshirani, \u201cRegression shrinkage and selection via the lasso,\u201d J. Roy. Statist. Soc. B, vol. 58, no. 1, pp. 267\u2013288, 1996.",
 "cortes": "C. Cortes and V. Vapnik, \u201cSupport-vector networks,\u201d Mach. Learn., vol. 20, no. 3, pp. 273\u2013297, 1995.",
 "friedman": "J. H. Friedman, \u201cGreedy function approximation: A gradient boosting machine,\u201d Ann. Statist., vol. 29, no. 5, pp. 1189\u20131232, 2001.",
 "kohavi": "R. Kohavi, \u201cA study of cross-validation and bootstrap for accuracy estimation and model selection,\u201d in Proc. Int. Joint Conf. Artif. Intell. (IJCAI), 1995, pp. 1137\u20131143.",
 "fawcett": "T. Fawcett, \u201cAn introduction to ROC analysis,\u201d Pattern Recognit. Lett., vol. 27, no. 8, pp. 861\u2013874, 2006.",
}
ORDER = []


def cite(text):
    def rep(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        nums = []
        for k in keys:
            if k not in REFS:
                raise KeyError(k)
            if k not in ORDER:
                ORDER.append(k)
            nums.append(ORDER.index(k) + 1)
        nums.sort()
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            if j - i >= 2:
                out.append(f"[{nums[i]}]\u2013[{nums[j]}]")
            else:
                out.extend(f"[{n}]" for n in nums[i:j + 1])
            i = j + 1
        return ", ".join(out)
    return re.sub(r"(?<![_^])\{([a-z0-9_, ]+)\}", rep, text)


# =========================================================================== document plumbing
d = docx.Document(os.path.join(SP, "template_t.docx"))
body = d.element.body
els = list(body.iterchildren())
title_p, auth_sec_end, sec1_end, sec2_end, anchor = els[0], els[3], els[8], els[9], els[90]
for i, el in enumerate(els):
    if i == 1 or 4 <= i <= 7 or 10 <= i <= 89 or i == 91:
        body.remove(el)
for child in list(anchor):
    if child.tag != qn("w:pPr"):
        anchor.remove(child)
ppr = anchor.find(qn("w:pPr"))
for tag in ("w:pStyle", "w:numPr"):
    e = ppr.find(qn(tag))
    if e is not None:
        ppr.remove(e)
SECT_2COL = copy.deepcopy(ppr.find(qn("w:sectPr")))

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")


def add_runs(p, text, size=None, italic=None, bold=None, font=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        r_b, r_i, sub, sup = bold, italic, False, False
        if part.startswith("**") and part.endswith("**"):
            part, r_b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, r_i = part[1:-1], True
        elif part.startswith("_{"):
            part, sub = part[2:-1], True
        elif part.startswith("^{"):
            part, sup = part[2:-1], True
        r = p.add_run(part)
        if r_b is not None: r.bold = r_b
        if r_i is not None: r.italic = r_i
        if sub: r.font.subscript = True
        if sup: r.font.superscript = True
        if size: r.font.size = Pt(size)
        if font: r.font.name = font
    return p


def new_par(style=None):
    p = d.add_paragraph(style=style)
    anchor.addprevious(p._p)
    return p


LAST = [None]


def P(text, style="Body Text"):
    p = add_runs(new_par(style), cite(text))
    LAST[0] = p
    return p


def H1(text): return add_runs(new_par("Heading 1"), text)
def H2(text): return add_runs(new_par("Heading 2"), text)
def H5(text): return add_runs(new_par("Heading 5"), text)
def BUL(text): return P(text, "bullet list")


EQN = [0]


def EQ(linear):
    """Linear-format (UnicodeMath) equation; converted to a native Word equation by finalize.ps1."""
    EQN[0] += 1
    p = new_par("Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("\u00a7EQ\u00a7" + linear + f"#({EQN[0]})")
    return EQN[0]


FIGN = [0]


def _picture_par(path, width):
    p = new_par("Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(FIGS, path), width=Inches(width))
    return p


def FIG(path, caption, width=3.35):
    FIGN[0] += 1
    _picture_par(path, width)
    add_runs(new_par("figure caption"), cite(caption))
    return FIGN[0]


def _sect(cols):
    s = copy.deepcopy(SECT_2COL)
    c = s.find(qn("w:cols"))
    if cols == 1:
        for a in list(c.attrib):
            del c.attrib[a]
        c.set(qn("w:space"), "720")
    return s


def FIGW(path, caption, width=7.0):
    """Full-width figure: close the running 2-column section on the previous paragraph,
    place the figure in a 1-column continuous section, then resume two columns."""
    FIGN[0] += 1
    prev = anchor.getprevious()
    pPr = prev.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        prev.insert(0, pPr)
    pPr.append(_sect(2))
    _picture_par(path, width)
    cap = add_runs(new_par("figure caption"), cite(caption))
    cap.paragraph_format.space_after = Pt(8)
    cap._p.get_or_add_pPr().append(_sect(1))
    return FIGN[0]


def _borders(tbl, top=12, header=6, bottom=12):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side, sz in (("top", top), ("bottom", bottom)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    for side in ("left", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblPr.append(jc)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("left", 50), ("right", 50)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)


def _grid(tbl, widths):
    """python-docx sizes the grid to the full page width; pin it to the real column widths."""
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    tblPr = tbl._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(int(sum(widths) * 1440))); tw.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
             "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
    kids = sorted(tblPr, key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
    for k in kids:
        tblPr.append(k)


def _cell_border(cell, side, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders"); tcPr.append(tb)
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
    tb.append(e)


TABN = [0]


def TABLE(caption, header, rows, widths, align=None, group_rows=(), font=7.5, note=None):
    """Booktabs-style table: heavy top/bottom rules, light rule under the header, no vertical lines."""
    TABN[0] += 1
    cap = add_runs(new_par("table head"), caption)
    cap.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1 + len(rows) + (1 if note else 0), cols=len(header))
    anchor.addprevious(t._tbl)
    _borders(t)
    _grid(t, widths)
    t.autofit = False
    for ri, row in enumerate([header] + rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]
            cell.width = Inches(widths[ci])
            p = cell.paragraphs[0]
            p.style = d.styles["table copy"]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            a = (align or ["L"] + ["C"] * (len(header) - 1))[ci]
            p.alignment = {"L": WD_ALIGN_PARAGRAPH.LEFT, "C": WD_ALIGN_PARAGRAPH.CENTER, "R": WD_ALIGN_PARAGRAPH.RIGHT}[a]
            add_runs(p, cite(val), size=font, bold=True if ri == 0 else None,
                     font="Segoe UI Symbol" if val == "✓" else None)
            if ri == 0:
                _cell_border(cell, "bottom", 6)
            if ri - 1 in group_rows and ri > 1:
                _cell_border(cell, "top", 4)
            if ri < len(rows) or note:
                p.paragraph_format.keep_with_next = True
    if note:
        nc = t.rows[-1].cells[0].merge(t.rows[-1].cells[-1])
        nc.width = Inches(sum(widths))
        _cell_border(nc, "top", 6)
        npar = nc.paragraphs[0]
        npar.style = d.styles["table copy"]
        add_runs(npar, note, size=6.5, italic=True)
    sp = new_par("Normal"); sp.paragraph_format.space_after = Pt(2)
    return TABN[0]


def ALGO(title, lines):
    t = d.add_table(rows=2, cols=1)
    anchor.addprevious(t._tbl)
    _borders(t, top=12, bottom=12)
    _grid(t, [3.35])
    t.autofit = False
    hc, bc = t.rows[0].cells[0], t.rows[1].cells[0]
    hc.width = bc.width = Inches(3.35)
    _cell_border(hc, "bottom", 6)
    hp = hc.paragraphs[0]
    hp.style = d.styles["Normal"]
    add_runs(hp, title, size=8)
    first = True
    for num, txt, indent in lines:
        p = bc.paragraphs[0] if first else bc.add_paragraph()
        first = False
        p.style = d.styles["Normal"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.18 + 0.14 * indent)
        p.paragraph_format.first_line_indent = Inches(-0.18) if num else Inches(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        hp.paragraph_format.keep_with_next = True
        if num:
            r = p.add_run(f"{num}:\u2002")
            r.font.size = Pt(7.5)
        add_runs(p, txt, size=8)
    sp = new_par("Normal"); sp.paragraph_format.space_after = Pt(2)


# =========================================================================== title & authors
for child in list(title_p):
    if child.tag == qn("w:r"):
        title_p.remove(child)
docx.text.paragraph.Paragraph(title_p, d).add_run(
    "VERDANT: A CPU-Deployable Machine Learning Platform for Plant Health Diagnosis, Crop Planning and Grounded Advice in Urban Vertical Farming")


def author_block(p, lines):
    p.style = d.styles["Author"]
    p.paragraph_format.space_before = Pt(6)
    for i, (txt, ital) in enumerate(lines):
        r = p.add_run(txt)
        r.italic = ital
        r.font.size = Pt(11 if i == 0 else 9)
        if "[" in txt:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if i < len(lines) - 1:
            r.add_break()


AFF = [("School of Engineering", True), ("Anurag University", True), ("Hyderabad, India", False)]
body.remove(sec1_end); body.remove(sec2_end)
at = d.add_table(rows=2, cols=3)
auth_sec_end.addprevious(at._tbl)
at.autofit = False
jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); at._tbl.tblPr.append(jc)
for row in at.rows:
    for c in row.cells:
        c.width = Inches(2.3)
for n in (1, 2, 3):
    author_block(at.rows[0].cells[n - 1].paragraphs[0],
                 [(f"[Student Name {n}]", False), ("Dept. of [Your Department]", True)] + AFF + [(f"[roll-number-{n}]@anurag.edu.in", False)])
g = at.rows[1].cells[0].merge(at.rows[1].cells[2])
author_block(g.paragraphs[0], [("[Guide Name]", False), ("[Designation], Dept. of [Your Department]", True)] + AFF + [("[guide-email]@anurag.edu.in", False)])

# =========================================================================== abstract
add_runs(new_par("Abstract"), "*Abstract*\u2014Urban vertical farms can grow food close to consumers, but the small operators who run most of them rarely have access to plant pathologists, agronomists or the computing resources that commercial decision-support tools assume. This paper presents VERDANT, an open, CPU-deployable platform that integrates several machine-learning components behind a single web interface: two pretrained Vision Transformers for leaf-disease diagnosis and species identification, a retrieval-augmented advisor that grounds a Gemma 3 language model in a curated agronomy knowledge base, a physics-based agronomic engine for vapour-pressure-deficit and growth-index analysis, and four lightweight models trained from scratch for disease screening, crop recommendation, growth-cycle prediction and greenhouse climate forecasting. Every learned model is compared against classical baselines on held-out data, and the main design choices are probed with ablation and sensitivity studies. The custom models reach 83.6% accuracy (ROC AUC 0.917) for tomato late-blight screening, 87.9% for five-class crop recommendation, R\u00b2 = 0.856 for lettuce growth prediction and R\u00b2 = 0.891 (RMSE 0.98 \u00b0C) for one-hour-ahead temperature forecasting, a 34% error reduction over a persistence forecast. Dense MiniLM retrieval answers 79.4% of paraphrased grower queries correctly at rank one, more than twice the rate of TF-IDF and BM25, and indexing answers together with questions raises this to 92.6%. All components run on a laptop CPU with per-request latencies between 0.1 and 206 ms. The experiments also expose the limits of synthetic training data, which are discussed openly together with a roadmap for field validation.")
add_runs(new_par("Keywords"), "*Keywords*\u2014urban farming, vertical farming, plant disease detection, vision transformer, convolutional neural network, random forest, retrieval-augmented generation, vapour pressure deficit, decision support system")

# =========================================================================== I. Introduction
H1("Introduction")
P("More than half of the world\u2019s population already lives in cities, and the United Nations projects that this share will reach about 68% by 2050 {un2018}. Feeding these cities is becoming harder as arable land shrinks, fresh water becomes scarce and food travels ever longer distances before it is eaten {fao2017}. Vertical farming and other forms of controlled-environment agriculture (CEA) respond to this pressure by growing crops in stacked layers indoors, under artificial light and with tightly managed nutrients and climate {despommier, kozai}. Because water is recirculated and space is used in three dimensions, such farms can produce considerably more food per square metre than open fields while using a fraction of the water {benke}.")
P("Running a vertical farm, however, is knowledge-intensive. A grower must recognise the first symptoms of disease, choose crops that suit the available conditions, hold electrical conductivity (EC), pH, light and vapour pressure deficit (VPD) inside narrow bands, and anticipate climate swings before they stress the plants. Pathogens and pests alone remove an estimated 10\u201340% of the global yield of staple crops {savary}, and in a closed growing room a single undetected infection can spread through a whole rack within days. Commercial operators can afford agronomists and proprietary climate software; the hobbyists, start-ups, schools and community gardens that make up much of urban agriculture usually cannot.")
P("Machine learning (ML) is widely studied as a way to close this expertise gap {liakos, kamilaris}. Convolutional networks classify leaf diseases from photographs with high accuracy {mohanty, ferentinos}, tree ensembles support crop and yield planning, Internet of Things (IoT) sensing makes continuous climate data cheap {wolfert, elijah}, and retrieval-augmented language models can answer questions from a curated knowledge source {lewis}. In practice these capabilities are published as separate prototypes, are rarely compared with simple baselines, and often assume a GPU or a paid cloud service. A grower who wants to diagnose a leaf, pick a crop, estimate a harvest date and ask a follow-up question must stitch several tools together.")
P("This paper presents VERDANT, a single platform that brings these functions together and evaluates each of them rigorously. Its main contributions are:")
BUL("**An integrated, CPU-only architecture** that serves vision, tabular, physiological and language components through one REST interface, with lazy loading, concurrency control and graceful degradation when an external model is unavailable.")
BUL("**A grounded agronomic advisor** that combines dense MiniLM retrieval over a curated knowledge base with Gemma 3 generation, together with a 68-query paraphrase benchmark showing that dense retrieval more than doubles the rank-1 accuracy of TF-IDF and BM25.")
BUL("**A physics-based agronomic engine** that turns air temperature, humidity, light, EC and CO\u2082 into VPD, daily light integral (DLI) and an interpretable growth index based on Liebig\u2019s law of the minimum.")
BUL("**Four lightweight models trained from scratch**, each compared with three to six classical baselines and analysed with ablation, depth-sensitivity and feature-importance studies, all reproducible from a public training notebook and experiment scripts.")
BUL("**An honest account of what the evidence does and does not show**, including cases where simpler baselines win and the effect of synthetic training data.")
P("The rest of the paper is organised as follows. Section II reviews related work. Section III describes the system architecture. Section IV details the datasets, models and evaluation protocol. Section V reports and discusses the results, and Section VI concludes with directions for future work.")

# =========================================================================== II. Related work
H1("Related Work")
H2("Image-Based Plant Disease Detection")
P("Deep learning has transformed image recognition since AlexNet won the ImageNet challenge {krizhevsky, deng, lecun}. The PlantVillage dataset, with more than 54,000 labelled leaf images of 14 crops and 26 diseases {hughes}, made it possible to apply these networks to plant pathology. Mohanty *et al.* {mohanty} reported 99.35% held-out accuracy with GoogLeNet, while noting a sharp drop on field images. Sladojevic *et al.* {sladojevic} recognised 13 disease types with an average precision of 96.3%, Ferentinos {ferentinos} reached 99.53% with VGG on 87,848 images of 25 species, Brahimi *et al.* {brahimi} visualised the leaf regions that drive tomato-disease predictions, and Too *et al.* {too} found DenseNet fine-tuning both accurate and parameter-efficient. Residual networks {he} have since become a default backbone, and the Plant Pathology 2020 challenge {thapa} confirmed that photographs with complex backgrounds remain much harder than laboratory images.")
P("The Transformer {vaswani} has more recently been adapted to vision. The Vision Transformer (ViT) {dosovitskiy} treats an image as a sequence of patches processed by self-attention, and data-efficient training recipes have produced compact variants such as DeiT-Tiny with about 5.7 M parameters {touvron}. Borhani *et al.* {borhani} showed that ViT-based models suit automated plant-disease classification and analysed their accuracy\u2013speed trade-off. VERDANT deploys public ViT checkpoints for diagnosis and, in parallel, trains a compact CNN from scratch to study what can be learned from very little data.")
H2("Crop Planning, Growth and Climate Prediction")
P("For tabular agronomic data, classical ML remains competitive. In their review, Liakos *et al.* {liakos} found that tree ensembles and support vector machines dominate yield-prediction and crop-selection studies. Random Forests {breiman} are popular because they model non-linear interactions between soil nutrients and weather, resist overfitting through bootstrap aggregation, and expose feature importances that growers can interpret, while regularised linear models such as Ridge {hoerl} provide transparent baselines for growth estimation. In greenhouse and plant-factory settings, Shamshiri *et al.* {shamshiri} surveyed climate-control strategies; recurrent networks such as LSTMs {hochreiter} have been used to forecast indoor temperature and humidity, and simple persistence and seasonal-naive forecasts remain the reference that any forecaster must beat {hyndman}.")
H2("Physiological Models in Controlled Environments")
P("Plant growth in CEA is governed by well-understood physics. The saturation vapour pressure of air follows Tetens\u2019 empirical equation {tetens}, and the resulting VPD controls transpiration, stomatal opening and nutrient uptake {monteith}. The DLI summarises the photosynthetically active photon dose a crop receives each day {kozai}, and Liebig\u2019s law of the minimum states that growth is limited by the scarcest resource {liebig}. These relationships are rarely exposed to growers in an interactive, quantitative form; VERDANT encodes them in a lightweight engine that runs in the browser.")
H2("Retrieval and Language-Model Advisors")
P("Pretrained language models, beginning with BERT {devlin}, made natural-language question answering practical. Sentence-BERT {reimers} produces embeddings that can be compared with cosine similarity, and MiniLM {wang} distils such models so that they run quickly on a CPU. Dense retrievers outperform sparse methods such as TF-IDF {salton} and BM25 {robertson} when queries paraphrase the stored text {karpukhin}. Retrieval-augmented generation (RAG) {lewis} conditions a generator on retrieved passages, which curbs hallucination and allows knowledge to be updated without retraining, and open-weight models such as Gemma 3 {gemma} make the pattern affordable. Table I contrasts representative work with VERDANT: most prior systems address one task, whereas VERDANT combines vision, tabular prediction, physiological modelling and grounded advice in one deployable platform.")
TABLE("Capabilities of Representative Work Compared with VERDANT",
      ["Work", "Dis.", "Spe.", "Crop", "Clim.", "Adv.", "Base."],
      [["Mohanty *et al.* {mohanty}", "\u2713", "\u2013", "\u2013", "\u2013", "\u2013", "\u2713"],
       ["Ferentinos {ferentinos}", "\u2713", "\u2013", "\u2013", "\u2013", "\u2013", "\u2713"],
       ["Too *et al.* {too}", "\u2713", "\u2013", "\u2013", "\u2013", "\u2013", "\u2713"],
       ["Borhani *et al.* {borhani}", "\u2713", "\u2013", "\u2013", "\u2013", "\u2013", "\u2713"],
       ["Lewis *et al.* {lewis}", "\u2013", "\u2013", "\u2013", "\u2013", "\u2713", "\u2713"],
       ["**VERDANT (ours)**", "\u2713", "\u2713", "\u2713", "\u2713", "\u2713", "\u2713"]],
      [1.19, 0.36, 0.36, 0.36, 0.36, 0.36, 0.36], group_rows=(5,),
      note="Dis.: disease diagnosis; Spe.: species identification; Crop: crop or growth planning; Clim.: climate forecasting; Adv.: natural-language advice; Base.: comparison with baselines.")

# =========================================================================== III. Architecture
H1("System Architecture")
H2("Design Goals")
P("VERDANT was designed around four goals: (i) *accessibility*\u2014every component must run on a commodity laptop without a GPU; (ii) *graceful degradation*\u2014the platform must stay useful when the network or the external language model is unavailable; (iii) *modularity*\u2014each model sits behind its own endpoint so that it can be replaced independently; and (iv) *reproducibility*\u2014all custom models, datasets and experiments can be regenerated from fixed random seeds.")
H2("Three-Tier Architecture")
P("Fig. 1 shows the resulting three-tier design. The *presentation tier* is a single-page application written in React 19 and TypeScript, bundled with Vite and styled with Tailwind CSS. It hosts the Plant Doctor, Species Scanner, AI Advisor and Grow Planner, plus a set of client-side agronomic tools (growth simulator, VPD matrix, nutrient lab and dashboard) described in Section IV-E. The *application tier* is a FastAPI service exposing five REST endpoints. The *model tier* holds three pretrained Transformers downloaded from the Hugging Face Hub, the four models trained from scratch, the agronomy knowledge base with its cached embeddings, and an external Gemma 3 endpoint reached through the Ollama API.")
FIGW("fig_architecture.png", "Three-tier architecture of VERDANT. Solid arrows show request paths; dotted arrows show the fallback taken when the external language model is unavailable. The four models trained from scratch are evaluated offline in this study (Section V).")
H2("Application Tier and Concurrency")
P("*POST /api/diagnose* and *POST /api/identify* accept a multipart image and return the five most probable classes; *POST /api/advisor* returns a grounded answer with its topic, retrieval score and related questions; *POST /api/growplan* returns structured climate targets for a crop; and *GET /api/health* reports which models are loaded. Models are loaded lazily into an in-memory cache guarded by one lock per model, so that concurrent first requests never load the same weights twice, and a background thread started from the application\u2019s lifespan hook warms the vision models and pre-computes the knowledge-base embeddings while the server already accepts traffic. Language-model calls are issued asynchronously with httpx so that a slow external response never blocks image inference. Invalid input yields HTTP 400 with a descriptive message, internal failures yield HTTP 500, and secrets are read from an environment file rather than source code.")

# =========================================================================== IV. Methodology
H1("Methodology")
H2("Datasets and Preprocessing")
P("The pretrained Transformers were used as released. For the four custom models, labelled data from real vertical farms were not available to the project, so four datasets were generated programmatically from value ranges and relationships taken from horticultural guidance {kozai}. Following common practice for synthetic training data {nikolenko}, noise, missing values, corrupted files and class overlap were injected on purpose so that the models face realistic preprocessing and generalisation problems rather than trivially separable data. Every random process uses a fixed seed. Table II summarises the datasets and Fig. 2 visualises them.")
P("*Leaf images:* A PlantVillage-style set {hughes} with two classes, tomato healthy and tomato late blight, was synthesised at 128\u2009\u00d7\u2009128 pixels with 320 images per class. The classes have overlapping colour ranges; lesions appear on only about 65% of diseased leaves to mimic early infections, and about 30% of healthy leaves carry harmless spots. Gaussian sensor noise (\u03c3 = 34 grey levels) and JPEG compression are applied. Three zero-byte files were detected and removed during cleaning, leaving 637 images split 80/20 into 509 training and 128 validation images.")
P("*Crop recommendation:* 1,200 samples with seven features\u2014nitrogen (N), phosphorus (P), potassium (K), temperature, humidity, pH and rainfall\u2014and five labels (rice, maize, chickpea, cotton and coffee). Gaussian measurement noise was added and 31, 27 and 26 values were removed from N, temperature and humidity; these were imputed with the class-wise median and outliers were capped at 1.5 times the interquartile range. A stratified split gives 960 training and 240 validation samples, standardised with training statistics only.")
P("*Lettuce growth:* 600 hydroponic cycles described by DLI (11\u201319 mol\u2009m^{\u22122}\u2009d^{\u22121}), air temperature (18\u201324 \u00b0C), EC (0.9\u20132.1 mS/cm), pH (5.5\u20136.5) and water temperature (17\u201322 \u00b0C). Days to harvest fall with light, temperature and EC and carry Gaussian noise of \u03c3 = 1.05 days (range 12.3\u201325.7, mean 18.6). Note that pH and water temperature were deliberately given no effect, which provides a check on whether a model attributes importance correctly.")
P("*Greenhouse sensors:* ten days of temperature, humidity and CO\u2082 at 10-minute intervals (1,441 steps) with diurnal cycles and sensor noise. Lag features were built for temperature at *t*\u22121, *t*\u22122 and *t*\u22126 and for humidity and CO\u2082 at *t*\u22121; the target is temperature one hour ahead (*t*\u2009+\u20096). The series is split chronologically without shuffling into 1,143 training and 286 validation steps, so that no future information leaks into training.")
TABLE("Datasets Used to Train and Evaluate the Custom Models",
      ["Dataset", "Train / Val", "Inputs", "Target", "Injected defects"],
      [["Leaf images", "509 / 128", "128\u00d7128 RGB", "2 classes", "Noise, overlap, 3 corrupt files"],
       ["Crop recommendation", "960 / 240", "7 numeric", "5 crops", "Noise, 84 missing values"],
       ["Lettuce growth", "480 / 120", "5 numeric", "Days", "Noise, 2 irrelevant inputs"],
       ["Greenhouse sensors", "1,143 / 286", "8 incl. lags", "T at +1 h", "Noise, temporal order"]],
      [0.82, 0.6, 0.62, 0.5, 0.88])
FIGW("fig_datasets.png", "Overview of the four datasets: (a) validation leaf images, (b) Pearson correlation of the crop-recommendation features after cleaning, (c) distribution of lettuce cultivation times, and (d) the first three days of the greenhouse sensor streams.")
H2("Pretrained Vision Transformer Classifiers")
P("Both image endpoints use Vision Transformers {dosovitskiy} run through the Hugging Face Transformers library {wolf}. An image is resized to 224\u2009\u00d7\u2009224 pixels and cut into *N* = 196 patches of 16\u2009\u00d7\u200916 pixels; each flattened patch *x*_{p}^{i} is linearly projected by *E*, a learnable class token is prepended and positional embeddings *E*_{pos} are added:")
EQ("z_0=[x_\"cls\";x_p^1 E;x_p^2 E;\u2026;x_p^N E]+E_\"pos\"")
P("The sequence passes through *L* encoder layers of multi-head self-attention and two-layer perceptrons with residual connections, where each head computes")
EQ("\"Attention\"(Q,K,V)=\"softmax\"((Q K^\u22a4)/\u221a(d_k)) V")
P("with query, key and value projections *Q*, *K*, *V* and head dimension *d*_{k}. Our inspection of the deployed checkpoints corrects an assumption made in earlier project documentation. The disease checkpoint {hfdisease} is a compact ViT with 5.5 M parameters, comparable to DeiT-Tiny {touvron}, that predicts 13 classes: healthy and diseased conditions of corn, potato, rice and wheat plus an *invalid* class for non-leaf images; its authors report about 98% validation accuracy. The species checkpoint {hfspecies} is a ViT-Base/16 with 85.8 M parameters that distinguishes 52 medicinal and crop plants. For diagnosis, the top label is mapped to a curated treatment plan containing a severity level, immediate actions and prevention steps.")
H2("Foliar Disease CNN")
P("To study what a small network can learn from limited data, a four-block convolutional neural network (Fig. 3) was implemented in PyTorch {paszke} and trained from random initialisation. Each block applies a 3\u2009\u00d7\u20093 convolution, batch normalisation {ioffe}, ReLU and 2\u2009\u00d7\u20092 max pooling, halving the resolution while doubling the channels from 16 to 128. The resulting 128\u2009\u00d7\u20098\u2009\u00d7\u20098 map is flattened, regularised with dropout {srivastava} (*p* = 0.4) and classified by two fully connected layers. The network has 1,146,882 parameters, about five times fewer than the disease ViT and 75 times fewer than ViT-Base. It was trained for 15 epochs with Adam {kingma} (learning rate 10^{\u22123}, weight decay 10^{\u22124}, batch size 32) and flip and brightness augmentation {shorten}, minimising the cross-entropy")
EQ("\u2112=-1/B \u2211_(i=1)^B \u2211_(c=1)^C y_(i,c) log \u0177_(i,c)")
P("over mini-batches of size *B* with *C* = 2 classes, one-hot labels *y* and predicted probabilities \u0177.")
FIGW("fig_cnn_arch.png", "Architecture of the foliar disease CNN trained from scratch. Each block halves the spatial resolution and doubles the channel depth; feature-map sizes are given as channels \u00d7 height \u00d7 width.")
H2("Tabular Learners")
P("*Crop recommender:* a Random Forest {breiman} of 150 trees from scikit-learn {pedregosa}, with depth limited to 7 and at least three samples per leaf. Each split minimises the Gini impurity")
EQ("G(S)=1-\u2211_(c=1)^C p_c^2")
P("of the child nodes, where *p*_{c} is the share of class *c* in node *S*, and the prediction is the majority vote of the trees. The depth limit was chosen with the sensitivity analysis in Section V-B.")
P("*Growth predictor:* Ridge regression {hoerl} on the five standardised inputs, which shrinks coefficients to reduce variance:")
EQ("(w\u0302,b\u0302)=\"arg min\"_(w,b)\u2211_(i=1)^n (y_i-w^\u22a4 x_i-b)^2+\u03b1\u2016w\u2016^2")
P("with \u03b1 = 1.5, giving only six parameters whose signs and magnitudes can be read directly by a grower.")
P("*Climate forecaster:* because the series is short and the model must answer in milliseconds, forecasting is framed as supervised regression on lag features rather than with a recurrent network. A Random Forest regressor of *T* = 100 trees (depth \u2264 8, leaf size \u2265 3) averages the tree outputs *f*_{j}:")
EQ("\u0177_(t+6)=1/T \u2211_(j=1)^T f_j (x_t)")
H2("Agronomic Decision Engine")
P("The client-side tools rest on a small physiological model, implemented in TypeScript and therefore free of any server cost. Saturation vapour pressure (kPa) at temperature *T* (\u00b0C) follows Tetens {tetens}:")
EQ("e_s (T)=0.61078 exp((17.27 T)/(T+237.3))")
P("and VPD is computed against a leaf that runs 1.5 \u00b0C below air temperature under transpiration {monteith}:")
EQ("\"VPD\"=e_s (T-1.5)-\"RH\"/100 e_s (T)")
P("The DLI follows from the photosynthetic photon flux density (PPFD, \u00b5mol\u2009m^{\u22122}\u2009s^{\u22121}) and the photoperiod *h* in hours, DLI = PPFD\u2009\u00b7\u2009*h*\u2009\u00b7\u20093600\u2009/\u200910^{6}. For each of eight crop profiles, every factor *k* (temperature, VPD, DLI, EC and CO\u2082) receives a Gaussian suitability")
EQ("s_k=exp(-1/2 ((v_k-v_k^*)/\u03c3_k )^2)")
P("around its optimum *v*_{k}^{*} with tolerance \u03c3_{k}; VPD uses a plateau over its healthy band and CO\u2082 a one-sided saturating score. Following Liebig\u2019s law of the minimum {liebig}, the growth index blends the weakest factor with the geometric mean, so that several mild stresses still compound:")
EQ("G=0.55 min_k s_k+0.45 (\u220f_(k=1)^5 s_k )^(1/5)")
P("The predicted cultivation time is the crop\u2019s base duration scaled by 1\u2009+\u20090.8(1\u2009\u2212\u2009*G*), and the factor with the lowest score is reported as the limiting factor. Fig. 4 visualises the VPD zones, the growth-index surface for lettuce and a worked example.")
FIGW("fig_agronomy.png", "Agronomic decision engine. (a) VPD zones over air temperature and relative humidity, from Tetens\u2019 equation with a 1.5 \u00b0C leaf offset. (b) Growth index *G* for butterhead lettuce over temperature and DLI (RH 65%, EC 1.2 mS/cm, CO\u2082 900 ppm); the star marks the optimum. (c) Factor scores for a stressed scenario (27 \u00b0C, 55% RH, DLI 14.4, EC 1.9, CO\u2082 600 ppm): temperature is the limiting factor and *G* falls to 0.50.")
H2("Retrieval-Augmented Advisor")
P("The advisor follows the RAG pattern {lewis} (Fig. 5 and Algorithm 1). Its knowledge base holds 34 curated question\u2013answer pairs across 12 topics, including nutrients, pH, lighting, pests, diseases, climate, water and hydroponic systems. At start-up every stored question is encoded by all-MiniLM-L6-v2 {reimers, wang}, a six-layer, 22.7 M-parameter encoder that produces 384-dimensional L2-normalised vectors. A user question *q* is encoded the same way and scored against each stored question *k*_{j} by cosine similarity:")
EQ("\"sim\"(q,k_j)=(q\u22c5k_j)/(\u2016q\u2016 \u2016k_j\u2016)")
P("The best-matching answer becomes the grounding context and is sent with a restrictive system prompt to Gemma 3 (4B) {gemma}. If no API key is configured, the call times out or the reply cannot be parsed, the retrieved answer is returned directly, so the advisor always produces a grounded response.")
FIG("fig_rag.png", "Retrieval-augmented advisor. Knowledge-base questions are embedded once at start-up; each request is embedded, matched by cosine similarity and answered by Gemma 3, with a deterministic fallback to the retrieved answer.")
ALGO("**Algorithm 1:** Grounded advisory inference",
     [(None, "**Input:** question *q*; knowledge base {(*k*_{j}, *a*_{j})}, *j* = 1\u202634; cached unit embeddings; threshold \u03c4 = 0.35", 0),
      (None, "**Output:** answer, topic, score *s*, related questions", 0),
      (1, "*e* \u2190 MiniLM(*q*) / \u2016MiniLM(*q*)\u2016", 0),
      (2, "*s*_{j} \u2190 *e* \u00b7 *k\u0302*_{j} for every *j*;  *m* \u2190 arg max_{j} *s*_{j}", 0),
      (3, "*ctx* \u2190 *a*_{m}", 0),
      (4, "**if** *s*_{m} < \u03c4 **then** prepend a low-confidence notice to *ctx*", 0),
      (5, "**if** an LLM key is configured **then**", 0),
      (6, "*ans* \u2190 Gemma3(system prompt, *ctx*, *q*)", 1),
      (7, "**on** timeout or parse error: *ans* \u2190 *ctx*", 1),
      (8, "**else** *ans* \u2190 *ctx*", 0),
      (9, "*rel* \u2190 up to three questions sharing the topic of entry *m*", 0),
      (10, "**return** (*ans*, topic(*m*), *s*_{m}, *rel*)", 0)])
H2("Experimental Protocol")
P("*Baselines:* each learned model is compared with classical alternatives trained on identical splits: for leaf images, colour-histogram features (48 histogram bins plus channel means and deviations) with logistic regression, an RBF support vector machine (SVM) {cortes} and a Random Forest; for crop recommendation, logistic regression, *k*-nearest neighbours, SVM, a decision tree, gradient boosting {friedman} and an unconstrained forest; for growth prediction, a mean predictor, ordinary least squares (OLS), Lasso {tibshirani}, *k*-NN, a Random Forest and gradient boosting; and for forecasting, persistence (\u0177_{t+6} = *T*_{t}), a 24-hour seasonal-naive forecast {hyndman}, linear regression, gradient boosting and the forest without CO\u2082 features. Five-fold cross-validation {kohavi} on the training split complements the single validation split.")
P("*Ablation:* the CNN was retrained without augmentation, without dropout, without batch normalisation, and without both augmentation and dropout, each with three random seeds; we report mean \u00b1 standard deviation.")
P("*Retrieval benchmark:* every knowledge-base entry received two paraphrased queries written to avoid the stored wording (68 queries in total). We report recall at rank *k* (R@*k*) and the mean reciprocal rank,")
EQ("\"MRR\"=(1)/(|Q|) \u2211_(i=1)^(|Q|) (1)/(\"rank\"_i)")
P("*Metrics and hardware:* classification uses accuracy, precision, recall, F1 and the area under the ROC curve {fawcett}; regression uses MAE,")
EQ("\"RMSE\"=\u221a(1/n \u2211_(i=1)^n (y_i-\u0177_i)^2)")
P("and the coefficient of determination")
EQ("R^2=1-(\u2211_i (y_i-\u0177_i)^2)/(\u2211_i (y_i-y\u0304)^2)")
P("Latency is the median over repeated single-input calls after warm-up on an AMD Ryzen 7 5700U laptop CPU with 15.4 GB RAM, Python 3.13, PyTorch 2.9.1, scikit-learn 1.6.1 and Transformers 5.4. Table III lists every model with its configuration and footprint. The saved models were re-evaluated by the experiment scripts and reproduced the notebook\u2019s metrics exactly.")
TABLE("Model Configurations and CPU Footprint",
      ["Model", "Configuration", "Params", "Latency"],
      [["Disease ViT\u2020", "ViT, 224\u00b2, patch 16, 13 classes", "5.5 M", "26.0 ms"],
       ["Species ViT\u2020", "ViT-B/16, 224\u00b2, 52 classes", "85.8 M", "205.9 ms"],
       ["MiniLM-L6-v2\u2020", "6 layers, 384-d, mean pooling", "22.7 M", "14.2 ms"],
       ["Foliar CNN", "4 conv blocks, Adam, 15 epochs", "1.15 M", "2.7 ms"],
       ["RF recommender", "150 trees, depth \u2264 7, leaf \u2265 3", "12.5 k nodes", "10.1 ms"],
       ["Ridge growth", "\u03b1 = 1.5, 5 inputs", "6", "0.08 ms"],
       ["RF forecaster", "100 trees, depth \u2264 8, leaf \u2265 3", "22.8 k nodes", "7.2 ms"]],
      [0.85, 1.38, 0.66, 0.5], align=["L", "L", "C", "C"], group_rows=(3,),
      note="\u2020 Pretrained model. Latency: median CPU time for one image or one query after warm-up.")

# =========================================================================== V. Results
H1("Results and Discussion")
H2("Foliar Disease Screening")
P("Fig. 6 summarises the CNN. Training and validation losses fall together over the first four epochs and then settle between 0.3 and 0.4, and the two accuracy curves stay within a few points of each other, so the network does not overfit despite only 509 training images. Validation accuracy peaked at 87.5% in epoch 13 and ended at 83.59% (weighted F1 = 0.84). The confusion matrix shows that 58 of 65 healthy and 49 of 63 blighted leaves were classified correctly; most errors are missed infections, giving a late-blight recall of 0.78 against 0.89 for healthy leaves. This matches the data: about 35% of diseased leaves carry no visible lesion. The ROC curve has an area of 0.917, so a lower decision threshold could trade some precision for higher recall on blight\u2014often the right trade in a closed growing room.")
FIGW("fig_cnn.png", "Foliar disease CNN: (a) training and validation loss, (b) accuracy, (c) validation confusion matrix with counts and row percentages, and (d) ROC curve of the deployed model.")
P("Table IV compares the CNN with classical baselines. A logistic-regression classifier on colour histograms reaches 89.84%, six points above the CNN. The ablation in Fig. 7 explains why. Removing brightness augmentation *raises* mean validation accuracy from 84.9% \u00b1 3.3 to 87.8% \u00b1 0.7, and removing augmentation and dropout together gives 87.5% \u00b1 1.1, whereas removing batch normalisation lowers it to 81.8% \u00b1 0.4. In this dataset the class signal lies mainly in global colour, and brightness jitter partly destroys it, so a model that sees raw colour statistics wins. Batch normalisation remains clearly beneficial, and dropout matters little at this scale. The lesson is that augmentation must preserve the cue that separates the classes; on real photographs, where lesion shape and texture carry the signal, the CNN\u2019s inductive bias is expected to matter much more {mohanty, too}.")
FIG("fig_ablation.png", "CNN ablation (mean \u00b1 std over three seeds). Removing brightness augmentation helps on this colour-driven dataset, while removing batch normalisation hurts; the dashed line marks the colour-histogram baseline.")
TABLE("Classification Baselines on the Validation Splits",
      ["Task / method", "Acc. (%)", "Macro-F1", "5-fold CV (%)"],
      [["*Leaf disease (2 classes)*", "", "", ""],
       ["Colour hist. + LogReg", "**89.84**", "**0.898**", "\u2013"],
       ["Colour hist. + SVM (RBF)", "88.28", "0.883", "\u2013"],
       ["Colour hist. + RF", "**89.84**", "**0.898**", "\u2013"],
       ["Foliar CNN (deployed)", "83.59", "0.836", "\u2013"],
       ["*Crop recommendation (5 classes)*", "", "", ""],
       ["Logistic regression", "**91.25**", "**0.910**", "**89.0 \u00b1 2.4**"],
       ["SVM (RBF)", "**91.25**", "0.908", "88.1 \u00b1 2.6"],
       ["*k*-NN (*k* = 5)", "87.50", "0.869", "85.4 \u00b1 1.6"],
       ["Decision tree (depth 7)", "83.75", "0.831", "81.3 \u00b1 2.3"],
       ["Gradient boosting", "89.58", "0.893", "85.3 \u00b1 2.5"],
       ["RF, unconstrained", "88.75", "0.884", "86.6 \u00b1 2.6"],
       ["RF, depth 7 (deployed)", "87.92", "0.873", "86.6 \u00b1 2.4"]],
      [1.42, 0.55, 0.55, 0.78], group_rows=(5,))
H2("Crop Recommendation")
P("The deployed forest reaches 93.02% training and 87.92% validation accuracy (weighted F1 = 0.876). Fig. 8 breaks this down. Chickpea is recognised perfectly because its low humidity and high potassium set it apart, while maize is hardest (recall 0.63), with 21% of maize samples predicted as cotton, whose nitrogen, temperature and humidity ranges overlap. Humidity (27.5%) and rainfall (24.9%) dominate the Gini importances, followed by nitrogen (16.3%) and potassium (14.8%), whereas temperature and pH contribute little because their ranges are similar for all five crops. The depth sweep in Fig. 8(d) shows that validation accuracy saturates near 88% from depth 6, while training accuracy keeps rising to 100% for an unconstrained forest; the chosen depth of 7 cuts the generalisation gap from 11.3 to 5.1 points at the same cross-validated accuracy (86.6%).")
FIGW("fig_crop.png", "Random Forest crop recommender: (a) row-normalised validation confusion matrix (%), (b) per-class precision, recall and F1, (c) Gini feature importance, and (d) train and validation accuracy as a function of maximum tree depth.")
P("Table IV also shows that the forest is not the most accurate model: logistic regression and the SVM reach 91.25%, and logistic regression has the best cross-validated accuracy (89.0% \u00b1 2.4). The synthetic classes are uniform boxes blurred by Gaussian noise, which linear boundaries separate well. The forest still offers per-feature explanations and robustness to unscaled or missing inputs, but on this evidence a linear model or a simple ensemble of the two would be a better default, and we report this rather than tune the comparison in our favour.")
H2("Growth Prediction and Climate Forecasting")
P("The Ridge model explains 85.6% of the variance in growing time (MAE 0.84 days, RMSE 1.03 days). Because the noise injected into the data has \u03c3 = 1.05 days, this RMSE is essentially the irreducible error, and Fig. 9(a\u2013b) confirms unbiased, approximately Gaussian residuals. The standardised coefficients recover the generating process: DLI (\u22122.11), air temperature (\u22120.85) and EC (\u22120.50) shorten the cycle, while pH (0.12) and water temperature (0.01) are correctly found to be nearly irrelevant. Table V shows that OLS and Lasso perform identically, as expected for a linear target, while flexible learners (forest, boosting, *k*-NN) do worse because they spend capacity on noise.")
P("The climate forecaster achieves R\u00b2 = 0.891, MAE 0.78 \u00b0C and RMSE 0.98 \u00b0C on the held-out final two days. It reduces RMSE by 34.2% relative to persistence and by 22.7% relative to the seasonal-naive forecast, and Fig. 9(c) shows that it tracks the diurnal cycle while smoothing sensor noise. Two observations temper this result. First, plain linear regression on the same lag features is marginally better (RMSE 0.95 \u00b0C), so the forest\u2019s non-linearity is not needed here. Second, the lagged CO\u2082 reading carries 79.7% of the forest\u2019s importance and removing all CO\u2082 features raises the RMSE to 1.15 \u00b0C: in the synthetic data CO\u2082 shares the daily cycle of temperature and acts as a clock. In a real greenhouse, where ventilation and enrichment decouple CO\u2082 from temperature, explicit time-of-day features should replace this proxy.")
FIGW("fig_regression.png", "Regression results: (a) predicted versus actual lettuce cultivation time with \u00b11 RMSE band, (b) residual distribution with Gaussian fit, and (c) one-hour-ahead temperature forecast against the observed value over the last two validation days.")
TABLE("Regression and Forecasting Baselines",
      ["Task / method", "MAE", "RMSE", "R\u00b2"],
      [["*Lettuce growth (days)*", "", "", ""],
       ["Mean predictor", "2.29", "2.73", "\u22120.010"],
       ["OLS", "0.84", "1.03", "0.856"],
       ["Lasso (\u03b1 = 0.05)", "**0.83**", "**1.02**", "**0.858**"],
       ["*k*-NN (*k* = 5)", "1.07", "1.30", "0.770"],
       ["Random Forest", "0.91", "1.13", "0.828"],
       ["Gradient boosting", "0.93", "1.13", "0.827"],
       ["Ridge (deployed)", "0.84", "1.03", "0.856"],
       ["*Greenhouse temperature, +1 h (\u00b0C)*", "", "", ""],
       ["Persistence", "1.19", "1.48", "0.749"],
       ["Seasonal naive (24 h)", "1.02", "1.26", "0.818"],
       ["Linear regression", "0.77", "**0.95**", "**0.898**"],
       ["Gradient boosting", "**0.77**", "0.96", "0.894"],
       ["RF without CO\u2082 lags", "0.92", "1.15", "0.848"],
       ["RF (deployed)", "0.78", "0.98", "0.891"]],
      [1.55, 0.55, 0.55, 0.6], group_rows=(8,))
H2("Advisor Retrieval")
P("Table VI and Fig. 10 report the paraphrase benchmark. The deployed MiniLM retriever ranks the correct entry first for 79.4% of queries (MRR 0.851), against 38.2% for TF-IDF and 35.3% for BM25 on the same question index, because paraphrases such as \u201cedges of my lettuce leaves are turning brown and crispy\u201d share almost no words with the stored \u201ctip burn\u201d question. Fourteen queries are missed. Most contain a term that appears only in the stored *answer* (for example \u201cpythium\u201d or \u201csubstrate\u201d) or a strong lexical anchor such as a crop name. A live test reproduced this failure mode: \u201cMy lettuce roots look brown and slimy in the DWC tank\u201d retrieved the lettuce tip-burn entry (score 0.61) instead of root rot. Indexing each question together with its answer addresses exactly this weakness and raises R@1 to 92.6% and R@3 to 98.5% at the same latency, so it will become the default. The score histogram in Fig. 10(b) shows that wrong matches concentrate below a cosine score of about 0.5, which supports the low-confidence notice at \u03c4 = 0.35 and suggests that a slightly higher threshold would catch more errors.")
TABLE("Retrieval Accuracy on 68 Paraphrased Queries",
      ["Retriever (index)", "R@1", "R@3", "MRR", "ms"],
      [["TF-IDF {salton} (Q)", "0.382", "0.647", "0.531", "0.4"],
       ["BM25 {robertson} (Q)", "0.353", "0.603", "0.505", "\u2013"],
       ["TF-IDF (Q + A)", "0.544", "0.779", "0.688", "0.4"],
       ["BM25 (Q + A)", "0.588", "0.853", "0.727", "\u2013"],
       ["MiniLM (Q), deployed", "0.794", "0.897", "0.851", "14.2"],
       ["MiniLM (Q + A)", "**0.926**", "**0.985**", "**0.959**", "14.8"]],
      [1.25, 0.42, 0.42, 0.42, 0.36], group_rows=(4,))
FIG("fig_retrieval.png", "Advisor retrieval: (a) R@1, R@3 and MRR for sparse and dense retrievers over question-only (Q) and question-plus-answer (Q+A) indices; (b) top-1 cosine scores of correct and incorrect matches for the deployed retriever, with the 0.35 confidence threshold.")
H2("Computational Footprint and System Demonstration")
P("Table III shows that every component fits comfortably on a laptop CPU. The heaviest, the 85.8 M-parameter species ViT, answers in about 206 ms; the disease ViT needs 26 ms and a full advisor retrieval 14 ms. The four custom models together occupy under 7.7 MB on disk and respond in 0.08\u201310 ms, so they could run on a single-board computer next to the sensors. Fig. 11 shows the deployed interface. The Plant Doctor returns the ranked classes with a treatment plan, the advisor shows its matched topic and retrieval score, and the Grow Simulator and VPD matrix expose the agronomic engine interactively.")
FIGW("fig_ui.png", "The deployed VERDANT interface (captured from the running system): (a) Plant Doctor diagnosing a potato leaf with the disease ViT, (b) AI Advisor answering a root-rot question through retrieval, (c) Grow Simulator showing the growth index and limiting factor, and (d) the VPD climate matrix.")
H2("Limitations and Threats to Validity")
P("Three limitations bound these conclusions. *External validity:* the four custom models were trained and evaluated on synthetic data. Their scores demonstrate a sound pipeline, not field performance, and the CNN results in particular show how a synthetic generator can reward shortcuts (here, global colour) that real images would not; retraining on photographs and sensor logs from a working farm is a precondition for deployment. *Pretrained components:* the accuracy figures for the ViT checkpoints are those reported by their authors and were not re-measured on vertical-farm imagery, whose lighting and backgrounds differ from public datasets {mohanty, thapa}; moreover, the disease ViT does not cover tomato, which is why the tomato-specific CNN was developed. *Advisor scope:* the knowledge base is small, the retrieval benchmark was written by the project team, and the quality of generated answers was not scored against a formal benchmark. Finally, the custom models are currently evaluated offline and are not yet exposed through the REST API.")

# =========================================================================== VI. Conclusion
H1("Conclusion and Future Work")
P("This paper presented VERDANT, a CPU-deployable platform that helps urban growers diagnose plant diseases, identify species, plan crops, estimate harvest dates, reason about climate and obtain grounded agronomic advice from one interface. Beyond integrating these functions, we evaluated every learned component against classical baselines and with ablation and sensitivity studies. Dense retrieval more than doubles the accuracy of keyword search for paraphrased questions, and a one-line change to the index raises it further to 92.6%. The forecaster cuts persistence error by a third. The baselines also show that simple models match or beat more complex ones on the current synthetic data.")
P("Future work will (i) collect real leaf images and sensor logs from a partner vertical farm to retrain and validate the custom models; (ii) expose the custom models through the API and connect live IoT sensors so the forecaster can raise alerts and drive actuators; (iii) replace the CO\u2082 proxy with explicit time features and compare against sequence models {hochreiter} on longer series; (iv) switch the advisor to question-plus-answer indexing, enlarge the knowledge base and evaluate generated answers with a formal benchmark; and (v) quantise the vision models for on-device inference in a mobile application.")
H5("Acknowledgment")
P("The authors thank their project guide, [Guide Name], and the School of Engineering, Anurag University, for their guidance and for providing the facilities needed to carry out this work.")
H5("References")
unused = [k for k in REFS if k not in ORDER]
if unused:
    print("UNCITED:", unused)
for k in ORDER:
    new_par("references").add_run(REFS[k])
d.save(OUT)
print("saved", OUT, "| refs:", len(ORDER), "| figs:", FIGN[0], "| tables:", TABN[0], "| eqs:", EQN[0])
